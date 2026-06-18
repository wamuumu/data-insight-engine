"""
Unit tests for src/ingestion/parsers/xlsx.py

Organised into:
  1. can_handle              — extension matching
  2. parse (end-to-end)      — happy path, multi-sheet warning, empty sheet
  3. _build_firmware_index   — column presence, event filtering
  4. _build_firmware_segments — all transition / anomaly cases documented in the code
  5. _make_firmware_lookup   — binary-search correctness, sentinel for out-of-range
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from common.constants import (
    SWITCH_OFF_EVENT_ID,
    SWITCH_ON_EVENT_ID,
    UNDEFINED_FIRMWARE_VERSION,
)
from ingestion.crawler.base import BaseFile
from ingestion.parsers.base import HistoryLogRecord
from ingestion.parsers.xlsx import XLSXParser


# ── helpers ───────────────────────────────────────────────────────────────────


def _make_file(tmp_path: Path) -> BaseFile:
    """Stub BaseFile pointing to a placeholder file (for index/segment tests)."""
    p = tmp_path / "device_B12345678.xlsx"
    p.write_bytes(b"placeholder")
    return BaseFile(p)


def _write_xlsx(
    path: Path,
    rows: list[tuple],
    header: list[str] | None = None,
    extra_sheets: int = 0,
) -> Path:
    header = header or ["Event ID", "Value", "Date", "Time"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(header)
    for row in rows:
        ws.append(list(row))
    for i in range(extra_sheets):
        wb.create_sheet(title=f"Extra{i}")
    wb.save(path)
    return path


# ── 1. can_handle ─────────────────────────────────────────────────────────────


class TestXLSXParserCanHandle:
    def test_true_for_xlsx(self, tmp_path: Path) -> None:
        p = tmp_path / "x.xlsx"
        p.write_bytes(b"stub")
        assert XLSXParser().can_handle(BaseFile(p)) is True

    def test_false_for_parquet(self, tmp_path: Path) -> None:
        p = tmp_path / "x.parquet"
        p.write_bytes(b"stub")
        assert XLSXParser().can_handle(BaseFile(p)) is False

    @pytest.mark.parametrize("ext", [".csv", ".txt", ".json", ""])
    def test_false_for_other_extensions(self, tmp_path: Path, ext: str) -> None:
        p = tmp_path / f"x{ext}"
        p.write_bytes(b"stub")
        assert XLSXParser().can_handle(BaseFile(p)) is False


# ── 2. parse (end-to-end) ─────────────────────────────────────────────────────


class TestXLSXParserParse:
    def test_yields_one_record_per_data_row(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[
                (SWITCH_ON_EVENT_ID, 101, "17/06/2026", "08:00:00.000"),
                (10, 5, "17/06/2026", "08:01:00.000"),
                (SWITCH_OFF_EVENT_ID, 202, "17/06/2026", "08:02:00.000"),
            ],
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        assert len(records) == 3

    def test_all_records_are_history_log_records(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(1, 100, "17/06/2026", "08:00:00.000")],
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        assert all(isinstance(r, HistoryLogRecord) for r in records)

    def test_firmware_version_propagated_within_on_off_segment(
        self, tmp_path: Path
    ) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[
                (SWITCH_ON_EVENT_ID, 101, "17/06/2026", "08:00:00.000"),
                (10, 5, "17/06/2026", "08:01:00.000"),
                (SWITCH_OFF_EVENT_ID, 202, "17/06/2026", "08:02:00.000"),
                (11, 6, "17/06/2026", "08:03:00.000"),
            ],
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        assert [r.data["firmware_version"] for r in records] == [101, 101, 101, 202]

    def test_undefined_firmware_when_no_switch_events(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(99, 55, "17/06/2026", "08:00:00.000")],
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        assert records[0].data["firmware_version"] == UNDEFINED_FIRMWARE_VERSION

    def test_required_fields_present_on_each_record(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(1, 100, "17/06/2026", "08:00:00.000")],
        )
        record = list(XLSXParser().parse(BaseFile(path)))[0]
        assert "event_date" in record.data
        assert "event_time" in record.data
        assert "event_id" in record.data
        assert "value" in record.data
        assert "firmware_version" in record.data

    def test_event_date_and_time_values_match_spreadsheet(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(1, 100, "25/12/2025", "23:59:59.999")],
        )
        record = list(XLSXParser().parse(BaseFile(path)))[0]
        assert record.data["event_date"] == "25/12/2025"
        assert record.data["event_time"] == "23:59:59.999"

    def test_source_file_is_path_string(self, tmp_path: Path) -> None:
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(1, 100, "17/06/2026", "08:00:00.000")],
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        assert records[0].source_file == str(path)

    def test_empty_sheet_produces_no_records(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "empty"
        path = tmp_path / "device_B12345678.xlsx"
        wb.save(path)
        records = list(XLSXParser().parse(BaseFile(path)))
        assert records == []

    def test_multi_sheet_file_still_parses(self, tmp_path: Path) -> None:
        """Extra sheets trigger a warning but should not crash the parser."""
        path = _write_xlsx(
            tmp_path / "device_B12345678.xlsx",
            rows=[(1, 100, "17/06/2026", "08:00:00.000")],
            extra_sheets=2,
        )
        records = list(XLSXParser().parse(BaseFile(path)))
        # Only Sheet1 has data rows; extra sheets are empty
        assert len(records) >= 1

    def test_corrupted_file_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "device_B12345678.xlsx"
        path.write_bytes(b"this is not xlsx")
        with pytest.raises(Exception):
            list(XLSXParser().parse(BaseFile(path)))

    def test_rows_with_all_none_are_skipped(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Event ID", "Value", "Date", "Time"])
        ws.append([1, 100, "17/06/2026", "08:00:00.000"])
        ws.append([None, None, None, None])  # blank row — should be skipped
        ws.append([2, 200, "17/06/2026", "08:01:00.000"])
        path = tmp_path / "device_B12345678.xlsx"
        wb.save(path)
        records = list(XLSXParser().parse(BaseFile(path)))
        assert len(records) == 2


# ── 3. _build_firmware_index ──────────────────────────────────────────────────


class TestBuildFirmwareIndex:
    def test_returns_empty_when_event_id_column_missing(self, tmp_path: Path) -> None:
        file = _make_file(tmp_path)
        result = XLSXParser()._build_firmware_index(file, ["Value"], [(100,)])
        assert result == []

    def test_returns_empty_when_value_column_missing(self, tmp_path: Path) -> None:
        file = _make_file(tmp_path)
        result = XLSXParser()._build_firmware_index(file, ["Event ID"], [(1,)])
        assert result == []

    def test_only_switch_events_are_returned(self, tmp_path: Path) -> None:
        file = _make_file(tmp_path)
        header = ["Event ID", "Value"]
        rows = [
            (SWITCH_ON_EVENT_ID, 100),
            (99, 999),  # non-switch — excluded
            (SWITCH_OFF_EVENT_ID, 200),
        ]
        events = XLSXParser()._build_firmware_index(file, header, rows)
        assert len(events) == 2
        assert events[0] == (0, SWITCH_ON_EVENT_ID, 100)
        assert events[1] == (2, SWITCH_OFF_EVENT_ID, 200)

    def test_non_numeric_cells_are_skipped(self, tmp_path: Path) -> None:
        file = _make_file(tmp_path)
        header = ["Event ID", "Value"]
        rows = [("N/A", "bad"), (SWITCH_ON_EVENT_ID, 100)]
        events = XLSXParser()._build_firmware_index(file, header, rows)
        assert len(events) == 1
        assert events[0][0] == 1  # second row (index 1)

    def test_all_non_switch_rows_returns_empty_list(self, tmp_path: Path) -> None:
        file = _make_file(tmp_path)
        header = ["Event ID", "Value"]
        rows = [(5, 100), (10, 200), (20, 300)]
        assert XLSXParser()._build_firmware_index(file, header, rows) == []


# ── 4. _build_firmware_segments ───────────────────────────────────────────────


class TestBuildFirmwareSegments:
    """
    Tests every transition documented in XLSXParser._build_firmware_segments.

    Notation: ON=SWITCH_ON_EVENT_ID, OFF=SWITCH_OFF_EVENT_ID
    """

    def _segments(self, switch_events, total_rows):
        parser = XLSXParser()
        tmp = Path("/tmp")
        file = _make_file(Path("/tmp"))
        return parser._build_firmware_segments(file, switch_events, total_rows)

    def test_no_events_no_segments(self, tmp_path) -> None:
        file = _make_file(tmp_path)
        segs = XLSXParser()._build_firmware_segments(file, [], total_rows=5)
        # With no switch events the trailing close still fires with UNDEFINED
        assert all(fw == UNDEFINED_FIRMWARE_VERSION for _, _, fw in segs)

    def test_simple_on_then_off(self, tmp_path) -> None:
        file = _make_file(tmp_path)
        events = [
            (0, SWITCH_ON_EVENT_ID, 100),
            (3, SWITCH_OFF_EVENT_ID, 200),
        ]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=5)
        # Rows 0-3 get firmware 100 (from ON); rows 4 onward get 200 (from OFF)
        start_vals = {s: fw for s, _, fw in segs}
        assert start_vals.get(0) == 100

    def test_orphan_off_at_start(self, tmp_path) -> None:
        """First event is OFF — rows before it (0..off_row) inherit its firmware."""
        file = _make_file(tmp_path)
        events = [(2, SWITCH_OFF_EVENT_ID, 555)]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=5)
        # First segment should be (0, 2, 555)
        assert segs[0] == (0, 2, 555)

    def test_on_on_anomaly_closes_first_segment(self, tmp_path) -> None:
        """ON → ON: first ON's segment closes at row before second ON."""
        file = _make_file(tmp_path)
        events = [
            (0, SWITCH_ON_EVENT_ID, 100),
            (3, SWITCH_ON_EVENT_ID, 200),
        ]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=5)
        # Segment for first ON should be (0, 2, 100)
        assert (0, 2, 100) in segs

    def test_off_off_anomaly_creates_two_separate_segments(self, tmp_path) -> None:
        """OFF → OFF: each OFF creates its own segment."""
        file = _make_file(tmp_path)
        events = [
            (1, SWITCH_OFF_EVENT_ID, 100),
            (3, SWITCH_OFF_EVENT_ID, 200),
        ]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=5)
        firmwares = [fw for _, _, fw in segs]
        assert 100 in firmwares
        assert 200 in firmwares

    def test_complex_sequence_matches_expected_segments(self, tmp_path) -> None:
        """Reproduce the exact scenario from the original test suite."""
        file = _make_file(tmp_path)
        events = [
            (0, SWITCH_ON_EVENT_ID, 100),
            (2, SWITCH_ON_EVENT_ID, 200),  # ON→ON anomaly
            (3, SWITCH_OFF_EVENT_ID, 300),
            (4, SWITCH_OFF_EVENT_ID, 400),  # OFF→OFF anomaly
            (5, SWITCH_ON_EVENT_ID, 500),
        ]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=6)
        assert (0, 1, 100) in segs
        assert (2, 3, 200) in segs
        assert (4, 4, 400) in segs
        assert (5, 5, 500) in segs

    def test_trailing_open_segment_closed_at_last_row(self, tmp_path) -> None:
        """If the last event is ON, the trailing segment should reach total_rows-1."""
        file = _make_file(tmp_path)
        events = [(2, SWITCH_ON_EVENT_ID, 777)]
        segs = XLSXParser()._build_firmware_segments(file, events, total_rows=5)
        # Trailing segment starts at 2, ends at 4 (last row index)
        assert any(end == 4 for _, end, _ in segs)

    def test_zero_rows_produces_no_segments(self, tmp_path) -> None:
        file = _make_file(tmp_path)
        segs = XLSXParser()._build_firmware_segments(
            file, [(0, SWITCH_ON_EVENT_ID, 1)], total_rows=0
        )
        assert segs == []


# ── 5. _make_firmware_lookup ──────────────────────────────────────────────────


class TestMakeFirmwareLookup:
    def _lookup(self, segments):
        return XLSXParser()._make_firmware_lookup(segments)

    def test_exact_segment_boundary_start(self) -> None:
        lookup = self._lookup([(2, 5, 42)])
        assert lookup(2) == 42

    def test_exact_segment_boundary_end(self) -> None:
        lookup = self._lookup([(2, 5, 42)])
        assert lookup(5) == 42

    def test_inside_segment_returns_firmware(self) -> None:
        lookup = self._lookup([(0, 9, 77)])
        for i in range(10):
            assert lookup(i) == 77

    def test_before_first_segment_returns_sentinel(self) -> None:
        lookup = self._lookup([(3, 6, 99)])
        assert lookup(0) == UNDEFINED_FIRMWARE_VERSION
        assert lookup(2) == UNDEFINED_FIRMWARE_VERSION

    def test_after_last_segment_returns_sentinel(self) -> None:
        lookup = self._lookup([(0, 4, 10)])
        assert lookup(5) == UNDEFINED_FIRMWARE_VERSION

    def test_gap_between_segments_returns_sentinel(self) -> None:
        lookup = self._lookup([(0, 2, 10), (5, 8, 20)])
        assert lookup(3) == UNDEFINED_FIRMWARE_VERSION
        assert lookup(4) == UNDEFINED_FIRMWARE_VERSION

    def test_multiple_adjacent_segments(self) -> None:
        lookup = self._lookup([(0, 2, 10), (3, 5, 20), (6, 8, 30)])
        assert [lookup(i) for i in range(9)] == [10, 10, 10, 20, 20, 20, 30, 30, 30]

    def test_empty_segments_always_returns_sentinel(self) -> None:
        lookup = self._lookup([])
        for i in range(5):
            assert lookup(i) == UNDEFINED_FIRMWARE_VERSION
